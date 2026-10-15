"""Peças usadas pelos três scripts da importação.

Regra que vale para tudo aqui: nenhuma função devolve nome ou CPF dentro de
mensagem de problema. Para apontar uma linha, usa-se a origem (arquivo, sítio,
tabela, linha), que é o que permite voltar ao papel.
"""
from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill

PASTA = Path(__file__).resolve().parent
TRABALHO = PASTA / "trabalho"

ABA_DOCUMENTOS = "documentos"
ABA_FORMATO_A = "formato_A"

# Colunas da aba formato_A, na ordem em que aparecem na planilha.
COLUNAS_ORIGEM = ["arquivo", "sitio", "tabela", "linha", "municipio", "uf", "data_referencia"]
# Texto cru do documento, sem interpretação (script 1).
COLUNAS_CRU = [
    "qtd_familia", "responsavel", "total_pessoas", "maior_60", "de_13_a_60",
    "crianca_ate_12_a", "crianca_ate_12_b", "gestante", "cpf",
]
# Preenchidas pelo script 2 só quando vazias; depois quem manda é a revisão.
COLUNAS_REVISAO = [
    "comunidade_nome", "comunidade_tipo", "comunidade_nova_confirmada",
    "meninos_ate_12", "meninas_ate_12", "acao",
]
# Recalculadas a cada execução do script 2 (e conferidas de novo no script 3).
COLUNAS_DIAGNOSTICO = ["cpf_situacao", "gravidade", "problemas"]
COLUNAS_A = COLUNAS_ORIGEM + COLUNAS_CRU + COLUNAS_REVISAO + COLUNAS_DIAGNOSTICO

COLUNAS_DOCUMENTOS = [
    "arquivo", "sha256", "formato", "tabelas", "sitio", "contato", "total_familias_cabecalho",
    "total_criancas_cabecalho", "linhas_numeradas_em_branco", "municipio", "uf",
    "data_referencia", "criado_em_metadado", "modificado_em_metadado",
]

# Títulos das colunas do documento (normalizados) -> coluna da planilha.
# "Criança até 12 anos" é mesclada em duas na grade: vira _a e _b.
CABECALHO_A = {
    "qtd familia": "qtd_familia",
    "responsavel": "responsavel",
    "total de pessoas": "total_pessoas",
    "maior de 60": "maior_60",
    ">13<60": "de_13_a_60",
    "crianca ate 12 anos": "crianca_ate_12",
    "gestante": "gestante",
    "cpf": "cpf",
}

DESCONHECIDO = {"", "?", "desconhecido", "desconhecida", "pendente"}

# Valores aceitos em "acao" (coluna de revisão).
ACAO_PULAR = "PULAR"
ACAO_CONFIRMADO = "CONFIRMADO"

# Prefixos que viram tipo de comunidade (enum TipoComunidade da API).
PREFIXOS_TIPO = {
    "sitio": "SITIO",
    "povoado": "POVOADO",
    "assentamento": "ASSENTAMENTO",
    "vila": "VILA",
    "distrito": "DISTRITO",
}

ERRO = "ERRO"
AVISO = "AVISO"


# ---------------------------------------------------------------- texto

def sem_acento(texto: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")


def normalizar(texto: str | None) -> str:
    """Chave de comparação: sem acento, minúsculo, pontuação vira espaço."""
    if not texto:
        return ""
    t = sem_acento(str(texto)).lower()
    t = re.sub(r"[^\w<>]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def texto(valor) -> str:
    """Célula da planilha como texto (openpyxl devolve None para vazio)."""
    if valor is None:
        return ""
    return str(valor).strip()


def eh_xx(valor: str) -> bool:
    return bool(re.fullmatch(r"X{2,}", valor.strip().upper()))


def separar_tipo_comunidade(nome_cabecalho: str) -> tuple[str, str]:
    """'SÍTIO GILA' -> ('Gila', 'SITIO'); 'SERROTA 01' -> ('Serrota 01', '')."""
    partes = nome_cabecalho.strip().split(maxsplit=1)
    tipo = ""
    if len(partes) == 2 and normalizar(partes[0]) in PREFIXOS_TIPO:
        tipo = PREFIXOS_TIPO[normalizar(partes[0])]
        nome = partes[1]
    else:
        nome = nome_cabecalho.strip()
    nome = " ".join(p.capitalize() if not p.isdigit() else p for p in nome.split())
    return nome, tipo


def chave_comunidade(nome: str) -> str:
    return normalizar(nome)


def chave_comunidade_fraca(nome: str) -> str:
    """Sem o prefixo de tipo: serve só para AVISAR de nome parecido, nunca para mesclar."""
    n = normalizar(nome)
    for p in list(PREFIXOS_TIPO) + ["comunidade", "assent", "sit"]:
        if n.startswith(p + " "):
            n = n[len(p) + 1:]
    return n


# ---------------------------------------------------------------- CPF

def cpf_digitos(valor: str) -> str:
    return re.sub(r"\D", "", valor or "")


def cpf_valido(digitos: str) -> bool:
    """Mesma regra de br.org...dominio.Cpf: 11 dígitos, não repetidos, DVs certos."""
    if not re.fullmatch(r"\d{11}", digitos) or len(set(digitos)) == 1:
        return False
    for pos in (9, 10):
        soma = sum(int(digitos[i]) * (pos + 1 - i) for i in range(pos))
        if (soma * 10) % 11 % 10 != int(digitos[pos]):
            return False
    return True


def situacao_cpf(valor: str) -> str:
    """Só a situação; o número nunca sai daqui."""
    v = valor.strip()
    if not v:
        return "AUSENTE"
    if eh_xx(v):
        return "XX"
    d = cpf_digitos(v)
    if len(d) != 11:
        return "MALFORMADO"
    if not cpf_valido(d):
        return "DV_INVALIDO"
    if not re.fullmatch(r"\d{11}|\d{3}\.\d{3}\.\d{3}-\d{2}", v):
        return "VALIDO_FORMATO_DIFERENTE"
    return "VALIDO"


# ---------------------------------------------------------------- números e crianças

def numero_inteiro(valor: str) -> int | None:
    v = valor.strip().rstrip("-").strip()
    return int(v) if re.fullmatch(r"\d{1,3}", v) else None


RE_NUMERO_SEXO = re.compile(r"^0*(\d{1,2})\s*([MF])$", re.IGNORECASE)


@dataclass
class CelulaCrianca:
    """Leitura de uma sub-célula de 'Criança até 12 anos'.

    tipo: VAZIO, XX, NUMERO_SEXO ('01 M'), NUMERO ('4'), OUTRO ('4 – 7', '1 A E 6M').
    NUMERO_SEXO não diz se o número é quantidade ou idade: isso é decisão humana.
    """
    tipo: str
    numero: int | None = None
    sexo: str | None = None


def ler_crianca(valor: str) -> CelulaCrianca:
    v = valor.strip()
    if not v:
        return CelulaCrianca("VAZIO")
    if eh_xx(v):
        return CelulaCrianca("XX")
    m = RE_NUMERO_SEXO.match(v)
    if m:
        return CelulaCrianca("NUMERO_SEXO", int(m.group(1)), m.group(2).upper())
    if re.fullmatch(r"\d{1,2}", v):
        return CelulaCrianca("NUMERO", int(v))
    return CelulaCrianca("OUTRO")


# ---------------------------------------------------------------- datas e origem

def ler_data(valor) -> date | None:
    if isinstance(valor, datetime):   # openpyxl devolve datetime para célula de data
        return valor.date()
    if isinstance(valor, date):
        return valor
    v = texto(valor)
    if normalizar(v) in DESCONHECIDO:
        return None
    try:
        return date.fromisoformat(v[:10])
    except ValueError:
        return None


def origem(linha: dict) -> str:
    """Como uma linha é citada em mensagem, log e relatório. Nunca nome, nunca CPF."""
    return f"{linha['arquivo']} · {linha['sitio']} · tabela {linha['tabela']} · linha {linha['linha']}"


def chave_origem(linha: dict) -> str:
    """Chave do arquivo de controle do envio."""
    return f"{linha['arquivo']}#t{linha['tabela']}l{linha['linha']}"


# ---------------------------------------------------------------- diagnóstico

@dataclass
class Problema:
    gravidade: str   # ERRO bloqueia o envio; AVISO não
    coluna: str
    texto: str


@dataclass
class Diagnostico:
    cpf_situacao: str = ""
    problemas: list[Problema] = field(default_factory=list)

    @property
    def gravidade(self) -> str:
        if any(p.gravidade == ERRO for p in self.problemas):
            return ERRO
        return AVISO if self.problemas else ""

    def add(self, gravidade: str, coluna: str, texto_: str) -> None:
        self.problemas.append(Problema(gravidade, coluna, texto_))


def propor_revisao(linha: dict) -> None:
    """Preenche as colunas de revisão que estiverem vazias. Não sobrescreve a revisão humana."""
    if not texto(linha.get("comunidade_nome")):
        nome, tipo = separar_tipo_comunidade(texto(linha["sitio"]))
        linha["comunidade_nome"] = nome
        if not texto(linha.get("comunidade_tipo")):
            linha["comunidade_tipo"] = tipo
    # Contagem por sexo só é proposta quando a linha inteira tem a forma "NN M/F"
    # e cabe no total. Mesmo assim o script 3 só usa com --criancas-como-contagem.
    if not texto(linha.get("meninos_ate_12")) and not texto(linha.get("meninas_ate_12")):
        celulas = [ler_crianca(texto(linha[c])) for c in ("crianca_ate_12_a", "crianca_ate_12_b")]
        uteis = [c for c in celulas if c.tipo not in ("VAZIO", "XX")]
        total = numero_inteiro(texto(linha["total_pessoas"]))
        if uteis and all(c.tipo == "NUMERO_SEXO" for c in uteis):
            m = sum(c.numero for c in uteis if c.sexo == "M")
            f = sum(c.numero for c in uteis if c.sexo == "F")
            if total is not None and m + f <= total - 1:
                linha["meninos_ate_12"] = m
                linha["meninas_ate_12"] = f


def diagnosticar(linhas: list[dict]) -> dict[int, Diagnostico]:
    """Problemas de cada linha (por índice). Usa a lista inteira para achar repetidos."""
    diags: dict[int, Diagnostico] = {i: Diagnostico() for i in range(len(linhas))}

    # repetidos: nome do responsável (entre e dentro de documentos) e CPF
    por_nome: dict[str, list[int]] = {}
    por_cpf: dict[str, list[int]] = {}
    for i, l in enumerate(linhas):
        if texto(l.get("acao")).upper() == ACAO_PULAR:
            continue
        n = normalizar(texto(l["responsavel"]))
        if n:
            por_nome.setdefault(n, []).append(i)
        d = cpf_digitos(texto(l["cpf"]))
        if cpf_valido(d):
            por_cpf.setdefault(d, []).append(i)

    for i, l in enumerate(linhas):
        dg = diags[i]
        acao = texto(l.get("acao")).upper()
        if acao and acao not in (ACAO_PULAR, ACAO_CONFIRMADO):
            dg.add(ERRO, "acao", f"valor de ação desconhecido: use vazio, {ACAO_PULAR} ou {ACAO_CONFIRMADO}")

        # origem do documento
        if normalizar(texto(l["municipio"])) in DESCONHECIDO:
            dg.add(ERRO, "municipio", "município não informado na extração (--municipio)")
        if not texto(l["uf"]) or not re.fullmatch(r"[A-Za-z]{2}", texto(l["uf"])):
            dg.add(ERRO, "uf", "UF ausente ou inválida")
        data_ref = ler_data(l["data_referencia"])
        if data_ref is None:
            dg.add(ERRO, "data_referencia", "data de referência do documento desconhecida: perguntar à associação")
        elif data_ref > date.today():
            dg.add(ERRO, "data_referencia", "data de referência no futuro")

        # responsável
        nome = texto(l["responsavel"])
        if not nome:
            dg.add(ERRO, "responsavel", "sem responsável: a API exige o nome do responsável da família")
        elif len(nome) > 120:
            dg.add(ERRO, "responsavel", "nome com mais de 120 caracteres (limite da API)")
        elif eh_xx(nome):
            dg.add(ERRO, "responsavel", "responsável marcado como XX")

        # total de pessoas
        bruto_total = texto(l["total_pessoas"])
        total = numero_inteiro(bruto_total)
        if total is None:
            if eh_xx(bruto_total):
                dg.add(AVISO, "total_pessoas", "total de pessoas XX: entra só o responsável; os demais ficam por cadastrar")
            elif not bruto_total:
                dg.add(AVISO, "total_pessoas", "total de pessoas em branco: entra só o responsável")
            else:
                dg.add(AVISO, "total_pessoas", f"total de pessoas não numérico ('{bruto_total}'): entra só o responsável")
        elif total == 0:
            dg.add(ERRO, "total_pessoas", "total de pessoas zero")
        elif total > 30:
            dg.add(ERRO, "total_pessoas", "total de pessoas acima de 30: conferir no papel")

        # faixas de adulto
        adultos = 0
        adultos_conhecidos = False
        for col in ("maior_60", "de_13_a_60"):
            v = texto(l[col])
            if not v:
                continue
            n = numero_inteiro(v)
            if n is None:
                dg.add(AVISO, col, f"valor não numérico ('{v}'): vai só para as observações")
            else:
                adultos += n
                adultos_conhecidos = True

        # crianças
        celulas = [ler_crianca(texto(l[c])) for c in ("crianca_ate_12_a", "crianca_ate_12_b")]
        for col, c in zip(("crianca_ate_12_a", "crianca_ate_12_b"), celulas):
            v = texto(l[col])
            if c.tipo == "OUTRO":
                dg.add(AVISO, col, f"formato não reconhecido ('{v}'): parece idade, não quantidade; vai só para as observações")
            elif c.tipo == "NUMERO":
                dg.add(AVISO, col, f"número sem sexo ('{v}'): quantidade ou idade? vai só para as observações")
            elif c.tipo == "NUMERO_SEXO" and c.numero > 12:
                dg.add(AVISO, col, f"'{v}': número acima de 12 numa coluna 'até 12 anos'")

        m_rev, f_rev = texto(l.get("meninos_ate_12")), texto(l.get("meninas_ate_12"))
        for col, v in (("meninos_ate_12", m_rev), ("meninas_ate_12", f_rev)):
            if v and numero_inteiro(v) is None:
                dg.add(ERRO, col, "deve ser um número (ou vazio)")
        uteis = [c for c in celulas if c.tipo not in ("VAZIO", "XX")]
        if uteis and total is not None:
            if all(c.tipo == "NUMERO_SEXO" for c in uteis):
                soma = sum(c.numero for c in uteis)
                limite = total - adultos if adultos_conhecidos else total - 1
                if (adultos_conhecidos and soma != limite) or (not adultos_conhecidos and soma > limite):
                    dg.add(AVISO, "crianca_ate_12_a",
                           f"lido como quantidade não fecha com o total ({soma} criança(s) para {total} pessoa(s)): provavelmente é idade")
        mr, fr = numero_inteiro(m_rev) or 0, numero_inteiro(f_rev) or 0
        if total is not None and mr + fr > total - 1:
            dg.add(ERRO, "meninos_ate_12", "meninos + meninas não cabem no total de pessoas (sobra o responsável)")

        if texto(l["gestante"]):
            dg.add(AVISO, "gestante", f"gestante preenchida ('{texto(l['gestante'])}') sem dizer quem: vai só para as observações")

        # CPF
        dg.cpf_situacao = situacao_cpf(texto(l["cpf"]))
        if dg.cpf_situacao in ("MALFORMADO", "DV_INVALIDO"):
            dg.add(ERRO, "cpf", "CPF inválido (a API recusa): corrigir pelo papel ou apagar a célula")
        elif dg.cpf_situacao in ("AUSENTE", "XX"):
            dg.add(AVISO, "cpf", "sem CPF no documento: família entra sem CPF")
        d = cpf_digitos(texto(l["cpf"]))
        if dg.cpf_situacao.startswith("VALIDO") and len(por_cpf.get(d, [])) > 1:
            outras = [origem(linhas[j]) for j in por_cpf[d] if j != i]
            dg.add(ERRO, "cpf", "mesmo CPF em outra linha (" + "; ".join(outras) + "): marcar PULAR em uma delas")

        n = normalizar(nome)
        if n and len(por_nome.get(n, [])) > 1 and acao != ACAO_CONFIRMADO:
            outras = [origem(linhas[j]) for j in por_nome[n] if j != i]
            dg.add(ERRO, "responsavel",
                   "responsável com o mesmo nome em outra linha (" + "; ".join(outras)
                   + "): possível família duplicada; PULAR uma ou CONFIRMADO se forem famílias diferentes")

        # comunidade
        if not texto(l.get("comunidade_nome")):
            dg.add(ERRO, "comunidade_nome", "comunidade sem nome")
        tipo = texto(l.get("comunidade_tipo"))
        if tipo and tipo not in ("SITIO", "ASSENTAMENTO", "POVOADO", "COMUNIDADE_QUILOMBOLA",
                                 "VILA", "DISTRITO", "BAIRRO", "OUTRO"):
            dg.add(ERRO, "comunidade_tipo", "tipo fora da lista fechada (TipoComunidade)")

    return diags


# ---------------------------------------------------------------- planilha

AMARELO = PatternFill("solid", fgColor="FFF2CC")
VERMELHO = PatternFill("solid", fgColor="F8CBAD")
CINZA = PatternFill("solid", fgColor="D9D9D9")
AZUL = PatternFill("solid", fgColor="DDEBF7")


def ler_aba(caminho: Path, aba: str) -> list[dict]:
    wb = load_workbook(caminho)
    if aba not in wb.sheetnames:
        return []
    ws = wb[aba]
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas:
        return []
    cab = [texto(c) for c in linhas[0]]
    return [dict(zip(cab, r)) for r in linhas[1:] if any(v not in (None, "") for v in r)]


def escrever_planilha(caminho: Path, documentos: list[dict], linhas_a: list[dict],
                      diags: dict[int, Diagnostico] | None = None) -> None:
    wb = Workbook()
    ws_doc = wb.active
    ws_doc.title = ABA_DOCUMENTOS
    _escrever(ws_doc, COLUNAS_DOCUMENTOS, documentos)

    ws = wb.create_sheet(ABA_FORMATO_A)
    _escrever(ws, COLUNAS_A, linhas_a)
    for col in COLUNAS_REVISAO:
        ws.cell(row=1, column=COLUNAS_A.index(col) + 1).fill = AZUL
    if diags:
        for i, dg in diags.items():
            r = i + 2
            for p in dg.problemas:
                if p.coluna in COLUNAS_A:
                    cel = ws.cell(row=r, column=COLUNAS_A.index(p.coluna) + 1)
                    cel.fill = VERMELHO if p.gravidade == ERRO else AMARELO
                    anterior = cel.comment.text + "\n" if cel.comment else ""
                    cel.comment = Comment(anterior + f"{p.gravidade}: {p.texto}", "validar.py")
            if dg.gravidade:
                ws.cell(row=r, column=COLUNAS_A.index("gravidade") + 1).fill = (
                    VERMELHO if dg.gravidade == ERRO else AMARELO)
    caminho.parent.mkdir(parents=True, exist_ok=True)
    tmp = caminho.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(caminho)


def _escrever(ws, colunas: list[str], linhas: list[dict]) -> None:
    ws.append(colunas)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = c.fill if c.fill and c.fill.fgColor.rgb not in (None, "00000000") else CINZA
    for l in linhas:
        ws.append([_celula(l.get(c)) for c in colunas])
    ws.freeze_panes = "A2"
    for i, c in enumerate(colunas, start=1):
        largura = max([len(c)] + [len(str(l.get(c) or "")) for l in linhas[:200]])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(largura + 2, 8), 60)
    # texto cru precisa continuar texto: "02" não pode virar 2
    for row in ws.iter_rows(min_row=2):
        for cel in row:
            if isinstance(cel.value, str):
                cel.number_format = "@"


def _celula(v):
    if isinstance(v, date):
        return v.isoformat()
    return v
